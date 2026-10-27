"""Unit tests for PDF/DOCX extraction. Run: python -m pytest tests/test_resume_extraction.py"""
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from app.services.resume_extractor import ResumeExtractionError, extract_resume


def make_docx():
    output = BytesIO()
    document = Document()
    document.add_heading("Test Candidate", 0)
    document.add_paragraph("Python Developer")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "FastAPI, SQL"
    document.save(output)
    return output.getvalue()


def test_docx_paragraphs_and_tables():
    result = extract_resume("resume.docx", make_docx())
    assert "Test Candidate" in result["extracted_text"]
    assert "FastAPI" in result["extracted_text"]
    assert result["file_type"] == "docx"


def test_pdf_without_text_needs_ocr():
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(output)
    with pytest.raises(ResumeExtractionError, match="No selectable text"):
        extract_resume("resume.pdf", output.getvalue())


def test_unsupported_extension():
    with pytest.raises(ResumeExtractionError, match="Only .pdf"):
        extract_resume("resume.txt", b"name")


def test_fake_pdf():
    with pytest.raises(ResumeExtractionError, match="not a valid PDF"):
        extract_resume("resume.pdf", b"just text")


def test_empty_file():
    with pytest.raises(ResumeExtractionError, match="empty"):
        extract_resume("resume.docx", b"")


def test_oversized_file():
    with pytest.raises(ResumeExtractionError, match="5 MB"):
        extract_resume("resume.pdf", b"%PDF-" + b"A" * (5 * 1024 * 1024))
